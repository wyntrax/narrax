#!/usr/bin/env python3
"""Build the Narrax setup runbook DOCX from its Markdown source of truth.

Usage:
  python scripts/docs/build_setup_docx.py \
    --source docs/NARRAX_END_TO_END_SETUP.md \
    --output docs/NARRAX_END_TO_END_SETUP.docx

Requirements: python-docx (see docs/NARRAX_END_TO_END_SETUP.md tooling notes).
The Markdown is authoritative; the generated DOCX is a presentation artifact.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
NAVY = RGBColor(0x0B, 0x1F, 0x33)
ACCENT = RGBColor(0x0E, 0x6B, 0x6B)
MUTED = RGBColor(0x5B, 0x67, 0x70)
INK = RGBColor(0x1B, 0x1F, 0x24)
CODE_BG = "F2F4F6"
TABLE_HEADER = "0B1F33"
TABLE_ALT = "F4F7F8"
RULE = "D5DCE1"

INLINE = re.compile(
    r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\n]+\*)"
)
HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBERED = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def clean(text: str) -> str:
    return text.replace("\u00a0", " ").strip()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for tag, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_paragraph_border(paragraph, color: str = RULE, size: int = 6, side: str = "bottom") -> None:
    if isinstance(color, RGBColor):
        color = str(color)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), "4")
    edge.set(qn("w:color"), color)
    borders.append(edge)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    def field_run(node) -> None:
        run = paragraph.add_run()
        run._r.append(node)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        field_run(node)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0E6B6B")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, *, bold: bool = False, italic: bool = False, code: bool = False) -> None:
    text = clean(text)
    if not text:
        return
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            run.bold = bold or None
            run.italic = italic or None
            if code:
                _style_code_run(run)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = italic or None
            if code:
                _style_code_run(run)
        elif token.startswith("__") and token.endswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = italic or None
            if code:
                _style_code_run(run)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.bold = bold or None
            run.italic = italic or None
            _style_code_run(run)
        elif token.startswith("[") and "](" in token:
            label, url = token[1:-1].split("](", 1)
            if url.startswith("http://") or url.startswith("https://"):
                add_hyperlink(paragraph, url, label)
            else:
                run = paragraph.add_run(label)
                run.bold = bold or None
                run.italic = italic or None
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.bold = bold or None
            if code:
                _style_code_run(run)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold or None
        run.italic = italic or None
        if code:
            _style_code_run(run)


def _style_code_run(run) -> None:
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x8A, 0x2B, 0x2B)
    r_pr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F0F1F2")
    r_pr.append(shd)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.12

    for name, size, color, before, after in (
        ("Heading 1", 26, NAVY, 0, 10),
        ("Heading 2", 16, ACCENT, 18, 7),
        ("Heading 3", 12.5, NAVY, 12, 4),
        ("Heading 4", 11, MUTED, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1


def setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.text = ""
    header.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), alignment=2)
    left = header.add_run("NARRAX  ·  END-TO-END SETUP AND OPERATIONS RUNBOOK")
    left.font.size = Pt(7.5)
    left.font.bold = True
    left.font.color.rgb = MUTED
    header.add_run("\t")
    right = header.add_run("22 October 2026")
    right.font.size = Pt(7.5)
    right.font.color.rgb = MUTED
    set_paragraph_border(header, RULE, 4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Narrax runbook  ·  page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    add_field(footer, "PAGE", "1")
    run = footer.add_run(" of ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    add_field(footer, "NUMPAGES", "1")


def add_cover(doc: Document, source: Path) -> None:
    for _ in range(4):
        doc.add_paragraph()
    wordmark = doc.add_paragraph()
    wordmark.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = wordmark.add_run("N A R R A X")
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    run.font.name = "Calibri"

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("End-to-End Setup\nand Operations Runbook")
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = NAVY
    run.font.name = "Calibri"

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Local workstation  →  GitHub Actions and GHCR  →  one Ubuntu VPS")
    run.font.size = Pt(14)
    run.font.color.rgb = MUTED

    rule = doc.add_paragraph()
    set_paragraph_border(rule, ACCENT, 18, "bottom")
    rule.paragraph_format.space_after = Pt(18)

    facts = [
        ("Prepared", "22 October 2026"),
        ("Audience", "Owner, deployment operator, future maintainer"),
        ("Repository", "narra-entitlements"),
        ("Product", "Narrax — Solana narrative, radar and smart-money intelligence"),
        ("Source of truth", source.as_posix()),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for label, value in facts:
        row = table.add_row()
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(12.4)
        p = row.cells[0].paragraphs[0]
        r = p.add_run(label.upper())
        r.font.size = Pt(7.5)
        r.font.bold = True
        r.font.color.rgb = MUTED
        p = row.cells[1].paragraphs[0]
        r = p.add_run(value)
        r.font.size = Pt(9.5)
        r.font.color.rgb = INK
        for cell in row.cells:
            set_cell_margins(cell, 40, 0, 40, 60)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    r = note.add_run(
        "This document documents the current implementation. It is not a claim that a production system is already running. "
        "The Markdown file named above is authoritative; this DOCX is generated from it. Rebuild after editing the source."
    )
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = MUTED

    doc.add_page_break()


def add_contents(doc: Document) -> None:
    heading = doc.add_paragraph("Contents", style="Heading 1")
    heading.paragraph_format.page_break_before = False
    note = doc.add_paragraph()
    r = note.add_run("Open in Word and choose “Update field” if page numbers are not yet populated.")
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = MUTED
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Table of contents — right-click and update field.")
    doc.add_page_break()


def add_code_block(doc: Document, lines: list[str]) -> None:
    if not lines:
        return
    paragraph = doc.add_paragraph()
    pf = paragraph.paragraph_format
    pf.left_indent = Cm(0.25)
    pf.right_indent = Cm(0.1)
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    set_paragraph_shading(paragraph, CODE_BG)
    for i, line in enumerate(lines):
        run = paragraph.add_run(line.replace("\t", "    "))
        run.font.name = "Consolas"
        run.font.size = Pt(8.2)
        run.font.color.rgb = RGBColor(0x21, 0x25, 0x2A)
        if i < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    columns = max(1, len(header))
    table = doc.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_row = table.rows[0]
    set_row_cant_split(header_row)
    set_repeat_table_header(header_row)
    for i, value in enumerate(header[:columns]):
        cell = header_row.cells[i]
        cell.text = ""
        set_cell_shading(cell, TABLE_HEADER)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, value.replace("`", ""), bold=True)
        for run in p.runs:
            run.font.size = Pt(8.6)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.bold = True
    for index, row_data in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for i in range(columns):
            value = row_data[i] if i < len(row_data) else ""
            cell = row.cells[i]
            cell.text = ""
            set_cell_margins(cell)
            if index % 2 == 1:
                set_cell_shading(cell, TABLE_ALT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(8.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int] | None:
    if start + 1 >= len(lines):
        return None
    head = lines[start].strip()
    sep = lines[start + 1].strip()
    if "|" not in head or not TABLE_SEP.match(sep):
        return None

    def cells(line: str) -> list[str]:
        line = line.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        return [clean(part) for part in line.split("|")]

    header = cells(head)
    rows: list[list[str]] = []
    i = start + 2
    while i < len(lines):
        line = lines[i]
        if not line.strip() or "|" not in line:
            break
        rows.append(cells(line))
        i += 1
    return header, rows, i


def render(doc: Document, markdown: str) -> tuple[int, int]:
    lines = markdown.splitlines()
    i = 0
    first_h1_seen = False
    tables = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            fence = stripped[3:].strip()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code)
            continue

        if not stripped:
            i += 1
            continue

        parsed = parse_table(lines, i)
        if parsed:
            header, rows, i = parsed
            add_table(doc, header, rows)
            tables += 1
            continue

        heading = HEADING.match(stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                # The cover already carries the document title; keep the Markdown title as a
                # compact running heading so the body still opens with the document identity.
                p = doc.add_paragraph(text, style="Heading 2")
                p.paragraph_format.page_break_before = False
            else:
                style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}[level]
                p = doc.add_paragraph(style=style)
                add_inline(p, text)
                if level == 1:
                    p.paragraph_format.page_break_before = True
                elif level == 2:
                    set_paragraph_border(p, RULE, 4)
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            set_paragraph_border(p, ACCENT, 12, "left")
            add_inline(p, stripped.lstrip("> ").strip(), italic=True)
            i += 1
            continue

        if re.match(r"^\s*(---+|\*\*\*+|___+)\s*$", stripped):
            p = doc.add_paragraph()
            set_paragraph_border(p, RULE, 6, "bottom")
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        bullet = BULLET.match(line)
        if bullet:
            indent = len(bullet.group(1).replace("\t", "  "))
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_inline(p, bullet.group(2))
            i += 1
            continue

        numbered = NUMBERED.match(line)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(3)
            prefix = p.add_run(f"{numbered.group(2)}.  ")
            prefix.font.bold = True
            add_inline(p, numbered.group(3))
            i += 1
            continue

        # Merge consecutive plain lines into one paragraph, as Markdown does for soft wraps.
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if not ns or HEADING.match(ns) or BULLET.match(nxt) or NUMBERED.match(nxt):
                break
            if ns.startswith("```") or ns.startswith(">") or ns.startswith("|"):
                break
            if re.match(r"^\s*(---+|\*\*\*+|___+)\s*$", ns):
                break
            paragraph_lines.append(ns)
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))
    return len(doc.paragraphs), tables


def build(source: Path, output: Path) -> None:
    markdown = source.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    setup_page(doc)
    add_cover(doc, source.relative_to(ROOT) if source.is_relative_to(ROOT) else source)

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    add_contents(doc)
    paragraphs, tables = render(doc, markdown)

    core = doc.core_properties
    core.title = "Narrax End-to-End Setup and Operations Runbook"
    core.subject = "Local workstation, GitHub Actions/GHCR, single-VPS production"
    core.author = "Narrax"
    core.keywords = "Narrax, Solana, deployment, runbook, GitHub Actions, VPS"
    core.comments = f"Generated from {source.name}; the Markdown source is authoritative."

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

    # Reopen to prove the archive is valid and report a few structural facts.
    check = Document(output)
    headings = sum(1 for p in check.paragraphs if p.style.name in {"Heading 1", "Heading 2", "Heading 3"})
    size_kb = output.stat().st_size / 1024
    print(f"Wrote {output} ({size_kb:.1f} KiB)")
    print(f"Paragraphs: {len(check.paragraphs)}; tables: {len(check.tables)}; headings: {headings}; source paragraphs rendered: {paragraphs}")
    if len(check.paragraphs) < 100 or len(check.tables) < 5:
        raise SystemExit("DOCX validation failed: too little content")
    if not any("Source references and document maintenance" in p.text for p in check.paragraphs):
        raise SystemExit("DOCX validation failed: final chapter missing")


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs/NARRAX_END_TO_END_SETUP.md")
    parser.add_argument("--output", type=Path, default=ROOT / "docs/NARRAX_END_TO_END_SETUP.docx")
    args = parser.parse_args(list(argv) if argv is not None else None)
    source = args.source if args.source.is_absolute() else (ROOT / args.source)
    output = args.output if args.output.is_absolute() else (ROOT / args.output)
    if not source.exists():
        print(f"Source not found: {source}", file=sys.stderr)
        return 2
    build(source, output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
